import base64
import io
from datetime import date, timedelta

from odoo import models, fields, api, _
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None


class TimesheetReportWizard(models.TransientModel):
    _name = 'timesheet.report.wizard'
    _description = 'Timesheet Report Wizard'

    project_id = fields.Many2one(
        'project.project',
        string='Project',
        help='Filter by a specific project. Leave empty for all projects.',
    )
    date_from = fields.Date(
        string='Date From',
        required=True,
        default=lambda self: self._default_date_from(),
    )
    date_to = fields.Date(
        string='Date To',
        required=True,
        default=lambda self: self._default_date_to(),
    )

    def _default_date_from(self):
        """Default to the start of the previous ISO week (last Monday)."""
        today = date.today()
        return today - timedelta(days=today.weekday() + 7)

    def _default_date_to(self):
        """Default to the end of the previous ISO week (last Sunday)."""
        today = date.today()
        return today - timedelta(days=today.weekday() + 1)

    def action_generate(self):
        """Generate the Excel report and return a download URL."""
        self.ensure_one()
        return self._generate_excel()

    def _get_entries(self):
        """Fetch timesheet entries matching the wizard criteria."""
        domain = [
            ('date', '>=', self.date_from),
            ('date', '<=', self.date_to),
        ]
        if self.project_id:
            domain.append(('project_id', '=', self.project_id.id))

        entries = self.env['timesheet.entry'].search(domain, order='date, id')
        if not entries:
            raise UserError(_(
                'No timesheet entries found for the selected date range.'
            ))
        return entries

    def _generate_excel(self):
        """Build the Excel workbook and return a download action."""
        if openpyxl is None:
            raise UserError(_(
                'The openpyxl library is required to generate Excel reports. '
                'Please install it with: pip install openpyxl'
            ))

        entries = self._get_entries()

        # ---- Styles ----
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font_white = Font(bold=True, size=11, color='FFFFFF')
        title_font = Font(bold=True, size=20)
        bold_font = Font(bold=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )
        center_align = Alignment(horizontal='center', vertical='center')

        wb = openpyxl.Workbook()

        # =====================================================================
        #  COVER SHEET
        # =====================================================================
        ws_cover = wb.active
        ws_cover.title = 'Cover'

        ws_cover.merge_cells('B2:F2')
        ws_cover['B2'] = 'Timesheet Report'
        ws_cover['B2'].font = title_font

        ws_cover['B4'] = 'Project:'
        ws_cover['B4'].font = bold_font
        ws_cover['C4'] = self.project_id.name if self.project_id else 'All Projects'
        ws_cover['B5'] = 'Date Range:'
        ws_cover['B5'].font = bold_font
        ws_cover['C5'] = f'{self.date_from}  to  {self.date_to}'
        ws_cover['B6'] = 'Generated:'
        ws_cover['B6'].font = bold_font
        ws_cover['C6'] = fields.Datetime.now().strftime('%Y-%m-%d %H:%M')

        # Logo placeholder
        ws_cover['B8'] = '[ Company Logo ]'
        ws_cover['B8'].font = Font(italic=True, color='999999')

        ws_cover.column_dimensions['B'].width = 18
        ws_cover.column_dimensions['C'].width = 40

        # =====================================================================
        #  EMPLOYEE SHEETS (one per employee with entries)
        # =====================================================================
        employees = entries.mapped('week_id.employee_id').sorted('name')

        for employee in employees:
            emp_entries = entries.filtered(
                lambda e: e.week_id.employee_id == employee
            ).sorted('date')

            # Sheet names limited to 31 chars
            sheet_name = employee.name[:31]
            ws = wb.create_sheet(title=sheet_name)

            # Title row
            ws.merge_cells('A1:D1')
            ws['A1'] = f'Timesheet Entries — {employee.name}'
            ws['A1'].font = Font(bold=True, size=14)

            # Column headers (row 3)
            headers = ['Date', 'Tasks', 'Description', 'Hours']
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=3, column=col_idx, value=header)
                cell.font = header_font_white
                cell.fill = header_fill
                cell.border = thin_border
                cell.alignment = center_align

            # Data rows — one row per DAY, with aggregated descriptions
            # Group entries by date
            from collections import defaultdict, OrderedDict
            by_date = OrderedDict()
            for entry in emp_entries:
                d = entry.date
                if d not in by_date:
                    by_date[d] = {'descriptions': [], 'hours': 0.0, 'tasks': set()}
                if entry.description:
                    by_date[d]['descriptions'].append(entry.description.strip())
                by_date[d]['hours'] += entry.hours
                if entry.task_id:
                    by_date[d]['tasks'].add(entry.task_id.name)

            row = 4
            running_total = 0.0
            for day_date, agg in by_date.items():
                day_name = day_date.strftime('%A')  # Monday, Tuesday, ...
                ws.cell(row=row, column=1, value=f'{day_name} {day_date.strftime("%Y-%m-%d")}')
                # List all tasks worked on that day
                ws.cell(row=row, column=2, value=', '.join(sorted(agg['tasks'])) if agg['tasks'] else '')
                # Aggregate description: join all descriptions with semicolons
                ws.cell(row=row, column=3, value='; '.join(agg['descriptions']) if agg['descriptions'] else '')
                ws.cell(row=row, column=4, value=agg['hours'])

                for col_idx in range(1, 5):
                    ws.cell(row=row, column=col_idx).border = thin_border

                running_total += agg['hours']
                row += 1

            # Total row
            ws.cell(row=row, column=1, value='TOTAL').font = bold_font
            ws.cell(row=row, column=4, value=running_total).font = bold_font
            for col_idx in range(1, 5):
                ws.cell(row=row, column=col_idx).border = thin_border

            # Column widths
            ws.column_dimensions['A'].width = 14
            ws.column_dimensions['B'].width = 34
            ws.column_dimensions['C'].width = 56
            ws.column_dimensions['D'].width = 12

        # =====================================================================
        #  SUMMARY SHEET
        # =====================================================================
        ws_summary = wb.create_sheet(title='Summary')

        ws_summary.merge_cells('A1:F1')
        ws_summary['A1'] = 'Timesheet Summary Report'
        ws_summary['A1'].font = Font(bold=True, size=16)

        # Find all weeks in the date range
        weeks = entries.mapped('week_id').sorted('week_start')

        # Headers: Employee | Total Hours | Week 1 | Week 2 | ...
        ws_summary['A3'] = 'Employee'
        ws_summary['B3'] = 'Total Hours'

        week_col_map = {}
        col = 3
        for week in weeks:
            col_letter = get_column_letter(col)
            ws_summary.cell(row=3, column=col, value=week.name or f'Week of {week.week_start}')
            week_col_map[week] = col
            col += 1

        # Style summary headers
        for c in range(1, col):
            cell = ws_summary.cell(row=3, column=c)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = center_align

        # Data rows
        row = 4
        for employee in employees:
            ws_summary.cell(row=row, column=1, value=employee.name)
            ws_summary.cell(row=row, column=1).border = thin_border

            emp_total = 0.0
            for week in weeks:
                week_hours = sum(
                    e.hours for e in entries
                    if e.week_id == week and e.week_id.employee_id == employee
                )
                if week_hours:
                    ws_summary.cell(row=row, column=week_col_map[week], value=week_hours)
                ws_summary.cell(row=row, column=week_col_map[week]).border = thin_border
                emp_total += week_hours

            ws_summary.cell(row=row, column=2, value=emp_total)
            ws_summary.cell(row=row, column=2).font = bold_font
            ws_summary.cell(row=row, column=2).border = thin_border

            row += 1

        # Grand total row
        ws_summary.cell(row=row, column=1, value='GRAND TOTAL').font = bold_font
        grand_total = sum(entry.hours for entry in entries)
        ws_summary.cell(row=row, column=2, value=grand_total).font = bold_font
        for c in range(1, col):
            ws_summary.cell(row=row, column=c).border = thin_border

        ws_summary.column_dimensions['A'].width = 30
        ws_summary.column_dimensions['B'].width = 14

        # =====================================================================
        #  SAVE & RETURN DOWNLOAD
        # =====================================================================
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        file_data = base64.b64encode(output.read())
        filename = f'timesheet_report_{self.date_from}_to_{self.date_to}.xlsx'

        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': file_data,
            'store_fname': filename,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }
